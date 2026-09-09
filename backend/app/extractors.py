import base64
from pathlib import Path

import pymupdf
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from openai import OpenAI

from .settings import settings


def extract_pdf(path):
    # First try normal PDF text extraction.
    text = "\n\n".join(
        p.extract_text() or ""
        for p in PdfReader(path).pages
    ).strip()

    if text:
        return text

    # Scanned/image-only PDF fallback:
    # render each page to an image and send images to OpenRouter vision.
    if not settings.openai_api_key:
        raise RuntimeError(
            "PDF contains no extractable text and OPENAI_API_KEY is not configured"
        )

    client = OpenAI(
        api_key=settings.openai_api_key,
        base_url=settings.openai_base_url,
    )

    doc = pymupdf.open(path)
    pages = []

    try:
        for i, page in enumerate(doc):
            pix = page.get_pixmap(
                matrix=pymupdf.Matrix(1.5, 1.5),
                alpha=False,
            )

            image_b64 = base64.b64encode(pix.tobytes("png")).decode("utf-8")

            pages.append(
                {
                    "type": "text",
                    "text": f"PAGE {i + 1}",
                }
            )

            pages.append(
                {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:image/png;base64,{image_b64}",
                    },
                }
            )
    finally:
        doc.close()

    response = client.chat.completions.create(
        model=settings.openai_model,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a document extraction system. "
                    "Extract all readable text from the supplied document pages. "
                    "Return only the document text. "
                    "Preserve headings, lists, tables, names, dates, numbers, "
                    "and structure as accurately as possible. "
                    "Do not summarize or invent content."
                ),
            },
            {
                "role": "user",
                "content": pages + [
                    {
                        "type": "text",
                        "text": "Transcribe all pages faithfully.",
                    }
                ],
            },
        ],
    )

    extracted = (response.choices[0].message.content or "").strip()

    if not extracted:
        raise RuntimeError("PDF OCR/text extraction returned no text")

    return extracted


def extract_docx(path):
    d = DocxDocument(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]

    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text.strip() for c in r.cells))

    return "\n".join(out).strip()


def extract_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []

    for ws in wb.worksheets:
        out.append(f"[SHEET: {ws.title}]")

        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]

            if any(vals):
                out.append(" | ".join(vals))

    return "\n".join(out).strip()


def extract(path, mime):
    ext = Path(path).suffix.lower()

    if mime == "application/pdf" or ext == ".pdf":
        return extract_pdf(path)

    if ext == ".docx":
        return extract_docx(path)

    if ext in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)

    if ext in {".txt", ".csv"}:
        return Path(path).read_text(errors="ignore")

    raise ValueError("Unsupported document type")
